import pandas as pd
from datetime import datetime, timedelta
from openpyxl import load_workbook
from openpyxl.styles import Border, Side, Font
from openpyxl.utils import get_column_letter
from fpdf import FPDF
import psycopg2
import os
import argparse
import requests
import time
import re


DB_CONFIG = {
    'host': 'localhost',
    'port': 5432,
    'dbname': 'db_name',
    'user': 'your_username',
    'password': 'your_password'
}

TELEGRAM_TOKEN = os.getenv('TELEGRAM_TOKEN')
CHAT_ID = os.getenv('CHAT_ID')

parser = argparse.ArgumentParser()
parser.add_argument("--report", choices=["daily", "weekly"], help="Тип отчета для теста")
args = parser.parse_args()

if args.report:
    report_type = args.report
else:
    report_type = "weekly" if datetime.now().weekday() == 0 else "daily"

now = datetime.now()
today = now.date()
yesterday = today - timedelta(days=1)

if report_type == "daily":
    output_date = (datetime.now().date() - timedelta(days=1)).strftime('%Y-%m-%d')
    formatted_date = (datetime.now().date() - timedelta(days=1)).strftime('%d.%m.%Y')
    output_file_xlsx = f"sla({output_date}).xlsx"
    query = f"""
    SELECT h.hostname,
           COALESCE((1 - SUM(d.downtime)::float / 3600 / 24 ) * 100, 100) AS sla
    FROM 
      (SELECT DISTINCT hostname 
       FROM public.m_downtime) h
    LEFT JOIN public.m_downtime d
      ON h.hostname = d.hostname
      AND d.time >= CURRENT_DATE - INTERVAL '1 day'
      AND d.time < CURRENT_DATE
    LEFT JOIN public.md_applications app
      ON h.hostname = app.hostname
    WHERE 
      app.is_test = 0  
      AND app.product != 'platform0'
    GROUP BY h.hostname
    ORDER BY sla ASC;
    """
elif report_type == "weekly":
    today = datetime.now().date()
    start_of_week = today - timedelta(days=today.weekday() + 7)  
    end_of_week = start_of_week + timedelta(days=6)  
    output_date = f"{start_of_week.strftime('%Y-%m-%d')}_{end_of_week.strftime('%Y-%m-%d')}"
    formatted_date = f"{start_of_week.strftime('%d.%m.%Y')} - {end_of_week.strftime('%d.%m.%Y')}"
    output_file_xlsx = f"sla({output_date}).xlsx"
    query = f"""
    SELECT h.hostname,
           COALESCE((1 - SUM(d.downtime)::float / 3600 / 24 / 7) * 100, 100) AS sla
    FROM 
      (SELECT DISTINCT hostname 
       FROM public.m_downtime) h
    LEFT JOIN public.m_downtime d
      ON h.hostname = d.hostname
      AND d.time >= '{start_of_week}' 
      AND d.time <= '{end_of_week}'
    LEFT JOIN public.md_applications app
      ON h.hostname = app.hostname
    WHERE 
      app.is_test = 0  
      AND app.product != 'platform0'
    GROUP BY h.hostname
    ORDER BY sla ASC;
    """

try:
    conn = psycopg2.connect(**DB_CONFIG)
    cursor = conn.cursor()
    cursor.execute(query)
    rows = cursor.fetchall()
    df = pd.DataFrame(rows, columns=["hostname", "sla"])
except Exception as e:
    print(f"Ошибка при подключении к БД или выполнении запроса: {e}")
    exit()
finally:
    if cursor:
        cursor.close()
    if conn:
        conn.close()

df.insert(0, "№", [i + 1 for i in range(len(df))])

avg_sla = df["sla"].mean()
df.loc[len(df)] = ["", "AVG", round(avg_sla, 4)]

base_folder = os.path.expanduser("~/sla_reporting")
month_folder = os.path.join(base_folder, now.strftime("%Y-%m"))
os.makedirs(month_folder, exist_ok=True)

hostnames_file = os.path.join(month_folder, f"hostnames_{report_type}.txt")

current_hostnames = set(df[df["hostname"] != "AVG"]["hostname"])
previous_hostnames = set()
if os.path.exists(hostnames_file):
    with open(hostnames_file, "r") as f:
        previous_hostnames = set(line.strip() for line in f)

with open(hostnames_file, "w") as f:
    for hostname in sorted(current_hostnames):
        f.write(f"{hostname}\n")

added = current_hostnames - previous_hostnames
removed = previous_hostnames - current_hostnames

change_summary = ""
if added or removed:
    change_summary += "\nИзменения в списке компаний:"
    if added:
        change_summary += f"\n➕ Добавлены: {', '.join(sorted(added))}"
    if removed:
        change_summary += f"\n➖ Отключены: {', '.join(sorted(removed))}"

print(change_summary)

excel_path = os.path.join(month_folder, output_file_xlsx)

with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
    df.to_excel(writer, index=False, sheet_name="SLA")
    wb = writer.book
    ws = wb.active
    
    thin_border = Border(left=Side(style='thin'),
                         right=Side(style='thin'),
                         top=Side(style='thin'),
                         bottom=Side(style='thin'))
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = thin_border
            cell.font = Font(name="Arial", size=10)
    
    ws.column_dimensions[get_column_letter(1)].width = 5  # Первая колонка (№)
    for col in range(2, ws.max_column + 1):
        max_length = max(len(str(cell.value)) for cell in ws[get_column_letter(col)])
        ws.column_dimensions[get_column_letter(col)].width = max_length + 2  # Запас



def send_document_with_caption(file_path, message):
    proxies = {
        "http": "http://192.168.8.2:3128",
        "https": "http://192.168.8.2:3128",
    }

    telegram_url = f"https://api.telegram.org/bot{TELEGRAM_TOKEN}/sendDocument"
    
    try:
        with open(file_path, "rb") as f:
            response = requests.post(
                telegram_url,
                data={
                    "chat_id": CHAT_ID,
                    "caption": message,
                    "parse_mode": "Markdown",
                },
                files={"document": f},
                proxies=proxies,
                timeout=15
            )
        
        if response.status_code == 200:
            print(f"[INFO] {now.strftime('%Y-%m-%d %H:%M:%S')} Документ с сообщением отправлен в Telegram", flush=True)
        else:
            print(f"[ERROR] Не удалось отправить сообщение: {response.status_code} {response.text}", flush=True)

    except Exception as e:
        print(f"[ERROR] Ошибка отправки документа в Telegram: {e}", flush=True)

final_message = f"Доброе утро!\nСредний SLA: {avg_sla:.4f}%\n{len(df)-1} компаний\n{formatted_date}{change_summary}"

send_document_with_caption(excel_path, final_message)

